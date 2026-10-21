"""Render the case document (``app.services.document``) as PDF, Word, HTML, Markdown or raw text,
and package several of them, with the evidence bundle, into one ZIP.

HTML and PDF share one template, so the PDF is the HTML printed. Word is built natively so
it stays editable: real headings, tables with repeating header rows, page numbers as
fields. All of them carry the classification on every page and the same reference.
"""

from __future__ import annotations

import base64
import hashlib
import io
import re
import zipfile
from datetime import UTC
from functools import cache
from pathlib import Path

from app.services.document import BASIS, Doc, clean

FORMATS = {
    "pdf": ("PDF", "application/pdf", "pdf"),
    "docx": ("Word", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "docx"),
    "html": ("HTML", "text/html; charset=utf-8", "html"),
    "md": ("Markdown", "text/markdown; charset=utf-8", "md"),
    "txt": ("Raw data (text)", "text/plain; charset=utf-8", "txt"),
    "evidence": ("Evidence bundle", "application/zip", "zip"),
}
FONTS = Path(__file__).resolve().parent.parent / "static" / "fonts"

INK, TEXT, MUTED, FAINT, RULE, TINT, MARK = "121110", "23221E", "57544C", "7A766C", "D6D2C7", "F6F5F0", "F2D34B"
SANS, MONO = "IBM Plex Sans", "IBM Plex Mono"


def slug(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-")[:60] or "case"


@cache
def font_css() -> str:
    faces = [("IBM Plex Sans", 400, "ibm-plex-sans-latin-400-normal"), ("IBM Plex Sans", 500,
             "ibm-plex-sans-latin-500-normal"), ("IBM Plex Sans", 600, "ibm-plex-sans-latin-600-normal"),
             ("IBM Plex Mono", 400, "ibm-plex-mono-latin-400-normal"), ("IBM Plex Mono", 500,
             "ibm-plex-mono-latin-500-normal")]  # fmt: skip
    out = []
    for family, weight, name in faces:
        data = base64.b64encode((FONTS / f"{name}.woff2").read_bytes()).decode()
        out.append(f'@font-face{{font-family:"{family}";font-weight:{weight};font-style:normal;'
                   f'src:url(data:font/woff2;base64,{data}) format("woff2");}}')  # fmt: skip
    return "\n".join(out)


# --- HTML and PDF -------------------------------------------------------------------------------


def render_html(doc: Doc) -> str:
    from app.web import templates

    return templates.env.get_template("exports/report.html").render(d=doc, font_css=font_css(), utc=UTC)


def render_pdf(doc: Doc) -> bytes:
    from weasyprint import HTML

    # Title, author and generator come from the document's <title> and <meta> tags.
    return HTML(string=render_html(doc), base_url=str(FONTS)).write_pdf()


# --- Markdown -----------------------------------------------------------------------------------


def _md(value: str) -> str:
    for ch in "\\`*_[]<>|#":
        value = value.replace(ch, "\\" + ch)
    return value


def _mdcode(value: str) -> str:
    value = value.replace("\n", " ").replace("|", "\\|")
    return f"`` {value} ``" if "`" in value else f"`{value}`"


def _date(dt) -> str:
    return dt.astimezone(UTC).strftime("%-d %b %Y") if dt else "—"


def _stamp(dt) -> str:
    return dt.astimezone(UTC).strftime("%-d %b %Y, %H:%M UTC") if dt else "—"


def render_markdown(doc: Doc) -> str:
    s = doc.stats
    kind = "Case brief" if doc.brief else "Case report"
    out = [
        f"**{doc.product}** · {kind} · **{doc.classification}**",
        "",
        f"# {_md(doc.title)}",
        "",
        "Targets: " + ", ".join(f"{t} {_mdcode(v)}" for t, v, _ in doc.targets),
        "",
        f"Prepared by {_md(doc.generated_by)} · {_stamp(doc.generated_at)} · ref `{doc.ref}` · {doc.scope_label}",
        "",
        f"> {doc.handling}",
        "",
        "## Summary",
        "",
        "### The analyst's conclusion",
        "",
        *[f"> {_md(p)}\n>" for p in doc.assessment],
        f"> *Last updated {_stamp(doc.assessment_updated_at)}.*",
        "",
        "| Confirmed | Likely | Possible | Connections | Sources |",
        "|---:|---:|---:|---:|---:|",
        f"| {s['confirmed']} | {s['likely']} | {s['possible']} | {s['links']} | {s['sources']} |",
        "",
        "### Strongest findings",
        "",
    ]
    out += [f"- **{f.strength_label}** · {f.type.replace('_', ' ')} {_mdcode(f.short)}" for f in doc.key] or [
        "Nothing has reached Likely or been confirmed yet."
    ]
    out += ["", "### Be careful about", ""]
    out += [f"- {_md(c)}" for c in doc.caveats] or ["Nothing flagged."]
    if doc.highlights:
        out += ["", "### Pinned by the analyst", ""]
        out += [f"- {_mdcode(w)}" + (f": {_md(y)}" if y else "") for w, y in doc.highlights]
    out += ["", "## Findings", "", f"{doc.scope_label}. Strength is how firmly a finding is tied to the subject; "
            "the grade (A–F) rates the kind of source.", ""]  # fmt: skip
    for label, rows, more in doc.groups:
        out += [f"### {label} ({len(rows) + more})", ""]
        out += ["| Finding | Strength | Score | Grade | Found by | First seen |", "|---|---|---:|:-:|---|---|"]
        for f in rows:
            extra = "; ".join(f"{k}: {v}" for k, v in f.details)
            flags = [x for x in (f.page_check, f"+{f.merged} merged" if f.merged else "",
                                 "imported, not observed" if f.imported else "") if x]  # fmt: skip
            cell = (
                _mdcode(f.value[:200])
                + (f"<br>{_md(extra)}" if extra else "")
                + (f"<br>*{_md(', '.join(flags))}*" if flags else "")
            )
            out.append(f"| {cell} | {f.strength_label} | {f.score:.2f} | {f.grade} | {_md(', '.join(f.sources))} "
                       f"| {_date(f.first_seen)} |")  # fmt: skip
        if more:
            out.append(f"\n*{more} weaker {label.lower()} not shown in the brief.*")
        out.append("")
    if doc.connections:
        out += ["## Connections", "", "Links by rule or similarity connect values, not people.", ""]
        for _, label, meaning, links in doc.connections:
            out += [f"### {label} ({len(links)})", "", f"*{meaning}*", ""]
            for c in links:
                arrow = "→" if c.directed else "↔"
                why = f" — {_md(c.why)}" if c.why else ""
                out.append(f"- {_mdcode(c.a)} {arrow} {_mdcode(c.b)}: {_md(c.label)}{why}")
            out.append("")
    if doc.contradictions:
        out += ["### Where sources disagree", ""] + [f"- **{_md(t)}.** {_md(x)}" for t, x in doc.contradictions] + [""]
    out += ["## Method & authorisation", "", "### Authorisation", ""]
    out += [f"> {_md(line)}" for line in doc.authorization] + [""]
    out += [f"- Lawful basis confirmed: {_stamp(doc.lawful_basis_at)}", f"- Case: `{doc.case.id}`",
            f"- Tools used: {', '.join(doc.tools) or '—'}",
            f"- Automatic follow-ups: {doc.pivots[0]} run, {doc.pivots[1]} skipped"]  # fmt: skip
    if doc.diff_summary:
        out.append(f"- Latest change: {doc.diff_summary}")
    out += ["", "### Scans", "", "| Scan | Started by | Status | When | Tools | Failed |", "|---|---|---|---|---:|---|"]
    out += [f"| #{r['number']} | {r['trigger']} | {r['status']} | {_stamp(r['at'])} | {r['tools']} | "
            f"{', '.join(r['failed']) or '—'} |" for r in doc.scans]  # fmt: skip
    if doc.gaps:
        out += ["", "### Coverage gaps", ""] + [f"- {g}" for g in doc.gaps]
    out += ["", "### How to read strength", ""] + [f"- **{k.replace('_', ' ').capitalize()}**: {v}"
                                                  for k, v in doc.strengths.items()]  # fmt: skip
    out += ["", "### Source grades", ""] + [f"- **{k}**: {v}" for k, v in doc.grades.items()]
    if not doc.brief:
        out += ["", "## Evidence", ""]
        if doc.snapshots:
            out += ["| Finding | Page | Captured | SHA-256 |", "|---|---|---|---|"]
            out += [f"| {_mdcode(x['finding'])} | {_md(x['url'])} | {_stamp(x['at'])} | `{x['sha256']}` |"
                    for x in doc.snapshots]  # fmt: skip
        else:
            out.append("No pages were captured for this case.")
    out += ["", "---", "", f"*End of {kind.lower()} · ref `{doc.ref}` · {doc.classification}. {doc.handling}*", ""]
    return "\n".join(out)


# --- Raw data -----------------------------------------------------------------------------------


def _raw(value) -> str:
    """Whole values, only control characters removed: this is the export for grep and diff."""
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]", "", str(value if value is not None else ""))


def render_text(doc: Doc) -> str:
    rule, thin = "=" * 78, "-" * 78
    out = [
        f"{doc.product} RAW CASE DATA   {doc.classification}",
        rule,
        f"Case:       {_raw(doc.case.name)}",
        f"Case id:    {doc.case.id}",
        f"Reference:  {doc.ref}",
        f"Generated:  {doc.generated_at.isoformat()} by {doc.generated_by}",
        f"Includes:   {doc.scope_label}",
        f"Handling:   {doc.handling}",
        "",
        "Targets",
        *[f"  {t:<10} {_raw(v)}" + (f"   [{', '.join(tags)}]" if tags else "") for t, v, tags in doc.targets],
        "",
        "Analyst assessment",
        *[f"  {_raw(p)}" for p in doc.assessment],
        "",
    ]
    findings = doc.all_findings
    out += [rule, f"FINDINGS ({len(findings)})", rule]
    for i, f in enumerate(findings, 1):
        out += [
            f"[{i}] {f.type}  {_raw(f.raw_value)}",
            f"    id:          {f.id}",
            f"    strength:    {f.strength_label} (score {f.score:.2f})",
            f"    grade:       {f.grade} ({f.grade_text})",
            f"    sources:     {', '.join(f.sources)}",
            f"    first seen:  {f.first_seen.isoformat()}",
            f"    last seen:   {f.last_seen.isoformat()}",
        ]
        if f.href:
            out.append(f"    url:         {f.href}")
        if f.page_check:
            out.append(f"    page check:  {f.page_check}")
        if f.merged:
            out.append(f"    merged:      {f.merged} duplicate value(s)")
        if f.imported:
            out.append("    imported:    yes, never observed by this platform")
        out += [f"    {k.lower() + ':':<13}{_raw(v)}" for k, v in f.details]
        out.append("")
    links = [(basis, c) for basis, _, _, cs in doc.connections for c in cs]
    out += [rule, f"CONNECTIONS ({len(links)})", rule]
    for basis, c in links:
        out.append(f"{_raw(c.a)} {'->' if c.directed else '<->'} {_raw(c.b)}")
        out.append(f"    {c.label}; basis {basis} ({BASIS[basis][0].lower()}); tool {c.tool}"
                   + (f"; confidence {c.confidence:.2f}" if c.confidence is not None else ""))  # fmt: skip
        if c.why:
            out.append(f"    why: {_raw(c.why)}")
    out += ["", rule, "SCANS", rule]
    out += [f"#{r['number']:<4} {r['trigger']:<11} {r['status']:<10} {r['at'].isoformat() if r['at'] else '-'}"
            f"  tools {r['tools']}" + (f"  failed: {', '.join(r['failed'])}" if r['failed'] else "")
            for r in doc.scans]  # fmt: skip
    if doc.contradictions or doc.gaps or doc.caveats:
        out += ["", rule, "CAVEATS", rule]
        out += [f"- {c}" for c in doc.caveats] + [f"- {t}: {x}" for t, x in doc.contradictions]
        out += [f"- {g}" for g in doc.gaps]
    if doc.snapshots:
        out += ["", rule, "CAPTURED PAGES", rule]
        for x in doc.snapshots:
            out += [f"{x['sha256']}  {x['at'].isoformat()}  HTTP {x['status']}", f"    {x['url']}",
                    f"    finding: {x['finding']}"]  # fmt: skip
    if doc.highlights:
        out += ["", rule, "PINNED ON THE WORKING LOG", rule] + [f"- {w}" + (f": {y}" if y else "")
                                                                for w, y in doc.highlights]  # fmt: skip
    out += ["", thin, f"Authorisation: {' '.join(_raw(x) for x in doc.authorization)}",
            f"Lawful basis confirmed: {doc.lawful_basis_at.isoformat() if doc.lawful_basis_at else '-'}",
            f"End of data. Reference {doc.ref}. {doc.classification}.", ""]  # fmt: skip
    return "\n".join(out)


# --- Word ---------------------------------------------------------------------------------------


@cache
def _mark_png() -> bytes:
    """The mark, drawn from its path on the 32-unit grid, in ink."""
    from PIL import Image, ImageDraw

    scale = 16
    img = Image.new("RGBA", (32 * scale, 32 * scale), (0, 0, 0, 0))
    draw = ImageDraw.Draw(img)
    pts = [(16, 26), (26, 16), (16, 6), (6, 26), (16, 16)]
    pts = [(x * scale, y * scale) for x, y in pts]
    w = round(4.4 * scale)
    draw.line(pts, fill=(18, 17, 16, 255), width=w, joint="curve")
    r = w / 2
    for x, y in (pts[0], pts[-1]):
        draw.ellipse((x - r, y - r, x + r, y + r), fill=(18, 17, 16, 255))
    buf = io.BytesIO()
    img.save(buf, "PNG")
    return buf.getvalue()


def render_docx(doc: Doc) -> bytes:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt, RGBColor

    def rgb(hexs: str) -> RGBColor:
        return RGBColor.from_string(hexs)

    def shade(el, fill: str) -> None:
        pr = el._tc.get_or_add_tcPr() if hasattr(el, "_tc") else el._r.get_or_add_rPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), fill)
        pr.append(shd)

    def borders(cell, **sides) -> None:
        pr = cell._tc.get_or_add_tcPr()
        b = OxmlElement("w:tcBorders")
        for side in ("top", "left", "bottom", "right"):
            e = OxmlElement(f"w:{side}")
            spec = sides.get(side)
            if spec:
                size, color = spec
                e.set(qn("w:val"), "single")
                e.set(qn("w:sz"), str(size))
                e.set(qn("w:color"), color)
            else:
                e.set(qn("w:val"), "nil")
            b.append(e)
        pr.append(b)

    def fixed(table, widths: list[float]) -> None:
        """Fixed column widths, so every word processor lays the table out the same way."""
        tbl_pr = table._tbl.tblPr
        layout = OxmlElement("w:tblLayout")
        layout.set(qn("w:type"), "fixed")
        tbl_pr.append(layout)
        grid = table._tbl.tblGrid
        for col, w in zip(grid.findall(qn("w:gridCol")), widths, strict=False):
            col.set(qn("w:w"), str(int(w * 56.7)))

    def no_borders(table) -> None:
        tbl_pr = table._tbl.tblPr
        b = OxmlElement("w:tblBorders")
        for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
            e = OxmlElement(f"w:{side}")
            e.set(qn("w:val"), "nil")
            b.append(e)
        tbl_pr.append(b)

    def font(run, name=SANS, size=None, bold=None, color=None, caps=False, spacing=None):
        run.font.name = name
        rpr = run._r.get_or_add_rPr()
        fonts = rpr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            rpr.append(fonts)
        for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            fonts.set(qn(attr), name)
        if size:
            run.font.size = Pt(size)
        if bold is not None:
            run.font.bold = bold
        if color:
            run.font.color.rgb = rgb(color)
        if caps:
            run.font.all_caps = True
        if spacing is not None:
            sp = OxmlElement("w:spacing")
            sp.set(qn("w:val"), str(spacing))
            rpr.append(sp)
        return run

    def para(container, text="", *, size=10, color=TEXT, name=SANS, bold=False, before=0, after=4, align=None,
             caps=False, spacing=None):  # fmt: skip
        p = container.add_paragraph()
        p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(before), Pt(after)
        if align:
            p.alignment = align
        if text:
            font(p.add_run(clean(text, 5000)), name, size, bold, color, caps, spacing)
        return p

    def eyebrow(container, audience: str):
        # Each part starts a page; a break set on the paragraph never leaves a blank page behind.
        p = para(container, after=2, before=0)
        p.paragraph_format.page_break_before = True
        font(p.add_run("FOR "), MONO, 7.5, False, MUTED, spacing=30)
        r = font(p.add_run(audience.upper()), MONO, 7.5, False, INK, spacing=30)
        shade(r, "FBF0B8")
        return p

    def heading(text: str, level: int):
        h = d.add_heading(level=level)
        h.paragraph_format.space_before = Pt(14 if level == 1 else 12)
        h.paragraph_format.space_after = Pt(6 if level == 1 else 4)
        h.paragraph_format.keep_with_next = True
        font(h.add_run(text), SANS, 17 if level == 1 else 11.5, True, INK)
        return h

    def field(paragraph, code: str):
        for kind, text in (("begin", None), (None, code), ("separate", None), (None, "1"), ("end", None)):
            r = paragraph.add_run()
            font(r, MONO, 7, False, FAINT)
            if kind:
                fc = OxmlElement("w:fldChar")
                fc.set(qn("w:fldCharType"), kind)
                r._r.append(fc)
            elif text == code:
                it = OxmlElement("w:instrText")
                it.set(qn("xml:space"), "preserve")
                it.text = f" {code} "
                r._r.append(it)
            else:
                r.text = text

    def tlp(paragraph, size=7.5):
        r = font(paragraph.add_run(f" {doc.classification} "), MONO, size, True, "FFC000", spacing=10)
        shade(r, "000000")
        return r

    def table(headers: list[str], widths: list[float]):
        t = d.add_table(rows=1, cols=len(headers))
        t.alignment = WD_TABLE_ALIGNMENT.LEFT
        t.autofit = False
        no_borders(t)
        fixed(t, widths)
        hdr = t.rows[0]
        tr_pr = hdr._tr.get_or_add_trPr()
        rep = OxmlElement("w:tblHeader")
        rep.set(qn("w:val"), "true")
        tr_pr.append(rep)
        for i, (h, w) in enumerate(zip(headers, widths, strict=True)):
            c = hdr.cells[i]
            c.width = Mm(w)
            borders(c, bottom=(12, INK))
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(3)
            font(p.add_run(h.upper()), MONO, 7, False, FAINT, spacing=16)
        return t

    def row(t, widths: list[float]):
        r = t.add_row()
        no_split = OxmlElement("w:cantSplit")
        no_split.set(qn("w:val"), "true")
        r._tr.get_or_add_trPr().append(no_split)
        for c, w in zip(r.cells, widths, strict=True):
            c.width = Mm(w)
            borders(c, bottom=(4, RULE))
        return r.cells

    def cell_text(c, text, *, name=SANS, size=8.8, color=TEXT, bold=False):
        p = c.paragraphs[0]
        p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(3), Pt(3)
        font(p.add_run(clean(text, 2000)), name, size, bold, color)
        return p

    def strength_run(p, f):
        r = font(p.add_run(f" {f.strength_label} "), MONO, 7.5, False, INK)
        if f.strength == "target":
            r.font.color.rgb = rgb("FFFFFF")
            shade(r, INK)
        elif f.strength == "confirmed":
            shade(r, MARK)
        elif f.strength in ("unlikely", "ruled_out"):
            r.font.color.rgb = rgb(FAINT)
            r.font.strike = f.strength == "ruled_out"
        return r

    d = Document()
    styles = d.styles
    normal = styles["Normal"]
    normal.font.name, normal.font.size = SANS, Pt(10)
    normal.font.color.rgb = rgb(TEXT)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), SANS)
    for s in ("Heading 1", "Heading 2", "Heading 3", "Title"):
        styles[s].font.name = SANS
        styles[s].font.color.rgb = rgb(INK)
    props = d.core_properties
    props.title = f"{doc.title} · {'Case brief' if doc.brief else 'Case report'}"
    props.author, props.subject = doc.generated_by, doc.classification
    props.keywords = f"{doc.product}; {doc.classification}; ref {doc.ref}"
    props.comments = doc.handling

    sec = d.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin = sec.right_margin = Mm(18)
    sec.top_margin, sec.bottom_margin = Mm(20), Mm(18)
    sec.header_distance = sec.footer_distance = Mm(9)
    sec.different_first_page_header_footer = True
    hp = sec.header.paragraphs[0]
    font(hp.add_run(f"{doc.product}  ·  {doc.title}"), MONO, 7, False, FAINT, spacing=12)
    hp.add_run("\t\t")
    tlp(hp, 7)
    fp = sec.footer.paragraphs[0]
    font(fp.add_run(f"Ref {doc.ref}  ·  {_stamp(doc.generated_at)}\t\tPage "), MONO, 7, False, FAINT)
    field(fp, "PAGE")
    font(fp.add_run(" of "), MONO, 7, False, FAINT)
    field(fp, "NUMPAGES")

    # Cover
    top = d.add_table(rows=1, cols=2)
    no_borders(top)
    left, right = top.rows[0].cells
    left.width, right.width = Mm(120), Mm(54)
    lp = left.paragraphs[0]
    lp.add_run().add_picture(io.BytesIO(_mark_png()), width=Mm(6.5))
    font(lp.add_run("  " + doc.product), MONO, 10.5, False, INK, spacing=48)
    rp = right.paragraphs[0]
    rp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    tlp(rp, 8.5)
    for c in (left, right):
        borders(c, bottom=(12, INK))
    para(d, after=0, before=60)
    para(d, "Case brief" if doc.brief else "Case report", name=MONO, size=9, color=MUTED, caps=True, spacing=36,
         after=8)  # fmt: skip
    tp = para(d, after=14)
    title = font(tp.add_run(doc.title), SANS, 28, True, INK)
    shade(title, "FBF0B8")
    tg = para(d, after=24)
    for i, (t, v, _tags) in enumerate(doc.targets):
        if i:
            font(tg.add_run("    "), SANS, 9)
        font(tg.add_run(t.upper() + " "), MONO, 7, False, FAINT, spacing=12)
        font(tg.add_run(v), MONO, 9.5, False, INK)
    meta = d.add_table(rows=2, cols=3)
    no_borders(meta)
    items = [("Prepared by", doc.generated_by), ("Generated", _stamp(doc.generated_at)), ("Reference", doc.ref),
             ("Includes", doc.scope_label), ("Findings", f"{doc.stats['findings']} across {doc.stats['scans']} scans"),
             ("Conclusion", f"Updated {_date(doc.assessment_updated_at)}")]  # fmt: skip
    for i, (k, v) in enumerate(items):
        c = meta.rows[i // 3].cells[i % 3]
        c.width = Mm(58)
        borders(c, top=(4, RULE) if i < 3 else None, bottom=(4, RULE) if i >= 3 else None)
        p = c.paragraphs[0]
        p.paragraph_format.space_before = Pt(6)
        font(p.add_run(k.upper()), MONO, 7, False, FAINT, spacing=20)
        p2 = c.add_paragraph()
        p2.paragraph_format.space_after = Pt(6)
        font(p2.add_run(v), MONO if k == "Reference" else SANS, 9.5, False, INK)
    para(d, after=0, before=90)
    how = d.add_table(rows=1, cols=3)
    no_borders(how)
    guide = [("Two minutes", "Read the Summary: the analyst's conclusion, the numbers and what to be careful about."),
             ("Investigating", "Findings and Connections list everything, how strongly each is believed, and why."),
             ("Checking the work", "Method & authorisation" + ("" if doc.brief else " and Evidence")
              + " show the basis for the case, how scores are meant and what was captured.")]  # fmt: skip
    for i, (h, txt) in enumerate(guide):
        c = how.rows[0].cells[i]
        c.width = Mm(58)
        borders(c, top=(24, MARK if i == 0 else INK))
        p = c.paragraphs[0]
        p.paragraph_format.space_before = Pt(6)
        font(p.add_run(h), SANS, 9.5, True, INK)
        font(c.add_paragraph().add_run(txt), SANS, 8.5, False, MUTED)
    hp2 = para(d, before=16)
    tlp(hp2)
    font(hp2.add_run("  " + doc.handling), SANS, 8, False, MUTED)

    # Summary
    eyebrow(d, "decision makers")
    heading("Summary", 1)
    box = d.add_table(rows=1, cols=1)
    no_borders(box)
    bc = box.rows[0].cells[0]
    bc.width = Mm(174)
    shade(bc, TINT)
    borders(bc, left=(36, MARK))
    p = bc.paragraphs[0]
    p.paragraph_format.space_before = Pt(6)
    font(p.add_run("THE ANALYST'S CONCLUSION"), MONO, 7.5, False, MUTED, spacing=20)
    for text in doc.assessment:
        q = bc.add_paragraph()
        q.paragraph_format.space_after = Pt(5)
        font(q.add_run(clean(text, 20000)), SANS, 11, False, INK)
    q = bc.add_paragraph()
    q.paragraph_format.space_after = Pt(6)
    font(q.add_run(f"Written by the analyst, last updated {_stamp(doc.assessment_updated_at)}. Scores below are "
                   "evidence for this conclusion, not the conclusion itself."), SANS, 8, False, FAINT)  # fmt: skip
    para(d, after=4)
    s = doc.stats
    st = d.add_table(rows=1, cols=5)
    no_borders(st)
    for i, (n, label) in enumerate([(s["confirmed"], "Confirmed"), (s["likely"], "Likely"),
                                    (s["possible"], "Possible"), (s["links"], "Connections"),
                                    (s["sources"], "Sources")]):  # fmt: skip
        c = st.rows[0].cells[i]
        c.width = Mm(34.8)
        borders(c, top=(8, "B9B3A5"), bottom=(8, "B9B3A5"), left=(4, RULE) if i else (8, "B9B3A5"),
                right=(8, "B9B3A5") if i == 4 else None)  # fmt: skip
        p = c.paragraphs[0]
        p.paragraph_format.space_before = Pt(6)
        font(p.add_run(str(n)), MONO, 20, False, INK)
        p2 = c.add_paragraph()
        p2.paragraph_format.space_after = Pt(6)
        font(p2.add_run(label), SANS, 8, False, MUTED)
    total = s["confirmed"] + s["likely"] + s["possible"] + s["unlikely"]
    if total:
        segs = [(k, s[k], c) for k, c in (("confirmed", INK), ("likely", "4A4740"), ("possible", "9C978B"),
                                          ("unlikely", RULE)) if s[k]]  # fmt: skip
        bar = d.add_table(rows=1, cols=len(segs))
        no_borders(bar)
        for c, (_k, n, color) in zip(bar.rows[0].cells, segs, strict=True):
            c.width = Mm(max(4.0, 174 * n / total))
            shade(c, color)
            c.paragraphs[0].paragraph_format.space_after = Pt(0)
            font(c.paragraphs[0].add_run(" "), SANS, 3)
        key = "   ".join(f"{k.capitalize()} {n}" for k, n, _ in segs)
        key += f"   Ruled out {s['ruled_out']}" if s["ruled_out"] else ""
        para(d, key, size=7.5, color=MUTED, before=3)
    heading("Strongest findings", 2)
    if doc.key:
        widths = [118, 56]
        t = table(["Finding", "Strength"], widths)
        for f in doc.key:
            c1, c2 = row(t, widths)
            p = c1.paragraphs[0]
            p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(3), Pt(3)
            font(p.add_run(f.type.replace("_", " ").upper() + "  "), MONO, 7, False, FAINT, spacing=10)
            font(p.add_run(f.short), MONO, 9, False, INK)
            p2 = c2.paragraphs[0]
            p2.paragraph_format.space_before = Pt(3)
            strength_run(p2, f)
            font(p2.add_run(f"  {len(f.sources)} source{'s' if len(f.sources) != 1 else ''}"), SANS, 8, False, FAINT)
    else:
        para(d, "Nothing has reached Likely or been confirmed yet.", color=FAINT, size=9)
    heading("Be careful about", 2)
    for c in doc.caveats or ["Nothing flagged: every tool completed, no duplicates are waiting, and no sources "
                             "contradict each other."]:  # fmt: skip
        p = d.add_paragraph(style="List Bullet")
        font(p.add_run(c), SANS, 9, False, TEXT)
    if doc.highlights:
        heading("Pinned by the analyst", 2)
        for w, y in doc.highlights[:6]:
            p = d.add_paragraph(style="List Bullet")
            font(p.add_run(w), MONO, 8.5, False, INK)
            if y:
                font(p.add_run(f" — {y}"), SANS, 8.5, False, MUTED)

    # Findings
    eyebrow(d, "investigators")
    heading("Findings", 1)
    para(d, f"Grouped by kind, strongest first. Strength is how firmly a finding is tied to the subject; the grade "
         f"rates the kind of source (A to F), independently of that. {doc.scope_label}.", size=10, color=MUTED,
         after=8)  # fmt: skip
    widths = [76, 30, 17, 30, 21]
    for label, rows, more in doc.groups:
        heading(f"{label}  ({len(rows) + more})", 2)
        t = table(["Finding", "Strength", "Grade", "Found by", "First seen"], widths)
        for f in rows:
            c = row(t, widths)
            p = c[0].paragraphs[0]
            p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(3), Pt(2)
            if f.host:
                font(p.add_run(f.host), MONO, 8.6, True, INK)
                font(p.add_run(f.path or ""), MONO, 8.6, False, INK)
            else:
                font(p.add_run(f.short), MONO, 8.6, False, INK)
            if f.details:
                font(c[0].add_paragraph().add_run(" · ".join(f"{k}: {v}" for k, v in f.details)), SANS, 7.8, False,
                     MUTED)  # fmt: skip
            flags = [x for x in (f.page_check, f"+{f.merged} merged" if f.merged else "",
                                 "Imported · not observed" if f.imported else "") if x]  # fmt: skip
            if flags:
                fl = c[0].add_paragraph()
                fl.paragraph_format.space_after = Pt(3)
                font(fl.add_run("  ".join(x.upper() for x in flags)), MONO, 6.8, False,
                     "8C5600" if f.imported else MUTED, spacing=8)  # fmt: skip
            p = c[1].paragraphs[0]
            p.paragraph_format.space_before = Pt(3)
            strength_run(p, f)
            font(c[1].add_paragraph().add_run(f"{f.score:.2f}"), MONO, 8, False, MUTED)
            cell_text(c[2], f.grade, name=MONO, color=INK)
            cell_text(c[3], ", ".join(f.sources), size=8, color=MUTED)
            cell_text(c[4], _date(f.first_seen), size=8, color=MUTED)
        if more:
            para(d, f"{more} weaker {label.lower()} not shown in the brief. The full report lists them.", size=8.5,
                 color=FAINT, before=3)  # fmt: skip
    if not doc.groups:
        para(d, "No findings in this export.", color=FAINT)

    if doc.connections:
        eyebrow(d, "investigators")
        heading("Connections", 1)
        para(d, "How findings relate to each other, grouped by how each link was established. Links by rule or "
             "similarity connect values, not people.", color=MUTED, after=8)  # fmt: skip
        for _, label, meaning, links in doc.connections:
            heading(f"{label}  ({len(links)})", 2)
            para(d, meaning, size=8.5, color=MUTED, after=4)
            for c in links:
                p = para(d, after=1, before=3)
                p.paragraph_format.keep_with_next = True
                font(p.add_run(c.a), MONO, 8.6, False, INK)
                font(p.add_run("  →  " if c.directed else "  ↔  "), SANS, 8.6, False, FAINT)
                font(p.add_run(c.b), MONO, 8.6, False, INK)
                extra = c.label + (f" · {c.why}" if c.why else "")
                para(d, extra, size=8, color=MUTED, after=3)
        if doc.contradictions:
            heading("Where sources disagree", 2)
            for t_, x in doc.contradictions:
                p = d.add_paragraph(style="List Bullet")
                font(p.add_run(t_ + ". "), SANS, 9, True, INK)
                font(p.add_run(x), SANS, 9, False, TEXT)

    # Method
    eyebrow(d, "legal, compliance and audit")
    heading("Method & authorisation", 1)
    heading("Authorisation", 2)
    q = d.add_table(rows=1, cols=1)
    no_borders(q)
    qc = q.rows[0].cells[0]
    qc.width = Mm(174)
    shade(qc, TINT)
    borders(qc, left=(18, "B9B3A5"))
    first = True
    for line in doc.authorization or ["—"]:
        p = qc.paragraphs[0] if first else qc.add_paragraph()
        first = False
        p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(3)
        font(p.add_run(clean(line, 5000)), SANS, 9.5, False, TEXT)
    kv = [("Lawful basis confirmed", _stamp(doc.lawful_basis_at)), ("Case reference", str(doc.case.id)),
          ("Tools used", ", ".join(doc.tools) or "—"),
          ("Automatic follow-ups", f"{doc.pivots[0]} run, {doc.pivots[1]} skipped")]  # fmt: skip
    if doc.diff_summary:
        kv.append(("Latest change", doc.diff_summary))
    para(d, after=2)
    t = d.add_table(rows=0, cols=2)
    no_borders(t)
    for k, v in kv:
        a, b = t.add_row().cells
        a.width, b.width = Mm(45), Mm(129)
        cell_text(a, k, size=9, color=MUTED)
        cell_text(b, v, size=9, color=INK, name=MONO if k == "Case reference" else SANS)
    heading("Scans", 2)
    widths = [14, 26, 24, 50, 16, 44]
    t = table(["Scan", "Started by", "Status", "When", "Tools", "Failed"], widths)
    for r in doc.scans:
        c = row(t, widths)
        cell_text(c[0], f"#{r['number']}", name=MONO)
        cell_text(c[1], r["trigger"])
        cell_text(c[2], r["status"])
        cell_text(c[3], _stamp(r["at"]), size=8, color=MUTED)
        cell_text(c[4], str(r["tools"]))
        cell_text(c[5], ", ".join(r["failed"]) or "—", size=8, color=MUTED)
    if doc.gaps:
        heading("Coverage gaps", 2)
        for g in doc.gaps:
            font(d.add_paragraph(style="List Bullet").add_run(g), SANS, 9, False, TEXT)
    heading("How to read strength", 2)
    t = d.add_table(rows=0, cols=2)
    no_borders(t)
    for k, v in doc.strengths.items():
        a, b = t.add_row().cells
        a.width, b.width = Mm(32), Mm(142)
        borders(a, bottom=(4, RULE))
        borders(b, bottom=(4, RULE))
        p = a.paragraphs[0]
        p.paragraph_format.space_before = Pt(3)

        class _F:  # the strength label styled as in the tables
            strength, strength_label = k, k.replace("_", " ").capitalize()

        strength_run(p, _F)
        cell_text(b, v, size=9)
    heading("Source grades", 2)
    t = d.add_table(rows=0, cols=2)
    no_borders(t)
    for k, v in doc.grades.items():
        a, b = t.add_row().cells
        a.width, b.width = Mm(32), Mm(142)
        borders(a, bottom=(4, RULE))
        borders(b, bottom=(4, RULE))
        cell_text(a, k, name=MONO, color=INK, bold=True)
        cell_text(b, v, size=9)
    heading("Limits", 2)
    para(d, "Scores combine the source's starting estimate, independent sources that agree, links between fields "
         "and matches to the case's context. Only identical identifiers are merged automatically; identity matches "
         "are decided by an analyst. A tool that could not complete is reported as failed, never as having found "
         "nothing. Open-source information can be wrong, out of date or planted: a finding is evidence about "
         "attribution, not proof of identity.", size=9.5)  # fmt: skip

    if not doc.brief:
        eyebrow(d, "legal, compliance and audit")
        heading("Evidence", 1)
        if doc.snapshots:
            para(d, "Pages captured during the case, exactly as received. The SHA-256 fingerprint identifies each "
                 "capture; the evidence bundle holds the bytes.", color=MUTED, after=8)  # fmt: skip
            widths = [38, 52, 28, 56]
            t = table(["Finding", "Page", "Captured", "SHA-256"], widths)
            for x in doc.snapshots:
                c = row(t, widths)
                cell_text(c[0], x["finding"], name=MONO, size=8)
                cell_text(c[1], f"{x['url']}  (HTTP {x['status']})", size=8)
                cell_text(c[2], _stamp(x["at"]), size=8, color=MUTED)
                cell_text(c[3], x["sha256"], name=MONO, size=6.8, color=MUTED)
        else:
            para(d, "No pages were captured for this case.", color=FAINT)
        if doc.highlights:
            heading("Pinned on the working log", 2)
            for w, y in doc.highlights:
                p = d.add_paragraph(style="List Bullet")
                font(p.add_run(w), MONO, 8.5, False, INK)
                if y:
                    font(p.add_run(f" — {y}"), SANS, 8.5, False, MUTED)

    end = para(d, before=24)
    font(end.add_run(f"End of {'brief' if doc.brief else 'report'} · Ref {doc.ref} · "), SANS, 8, False, MUTED)
    tlp(end)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


# --- Packaging ----------------------------------------------------------------------------------


def package(files: dict[str, bytes], doc: Doc) -> bytes:
    """Several exports in one ZIP, with a README listing each file's SHA-256."""
    lines = [f"{doc.product} export: {doc.title}", f"Reference {doc.ref} · {_stamp(doc.generated_at)} · "
             f"{doc.classification}", "", doc.handling, "", "Files (SHA-256):"]  # fmt: skip
    lines += [f"  {hashlib.sha256(body).hexdigest()}  {name}" for name, body in sorted(files.items())]
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("README.txt", "\n".join(lines) + "\n")
        for name, body in sorted(files.items()):
            z.writestr(name, body)
    return buf.getvalue()
